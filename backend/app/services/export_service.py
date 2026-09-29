import io
import csv
import json
from typing import List, Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExportService:
    def export_checklist_csv(self, checklist: List[Dict[str, Any]]) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Item ID", "Requirement", "Clause", "Evidence Needed", "Priority", "Completed", "Owner", "Due Date", "Notes"])
        for item in checklist:
            writer.writerow([
                item.get("id", ""),
                item.get("requirement", ""),
                item.get("clause", ""),
                item.get("evidence_needed", ""),
                item.get("priority", ""),
                "YES" if item.get("completed") else "NO",
                item.get("owner", ""),
                item.get("due_date", ""),
                item.get("notes", "")
            ])
        return output.getvalue().encode("utf-8")

    def export_checklist_json(self, checklist: List[Dict[str, Any]], metadata: Dict[str, Any]) -> bytes:
        data = {
            "title": "MANAK-AI Pre-Audit Compliance Checklist",
            "metadata": metadata,
            "checklist": checklist
        }
        return json.dumps(data, indent=2).encode("utf-8")

    def export_checklist_excel(self, checklist: List[Dict[str, Any]], metadata: Dict[str, Any]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pre-Audit Checklist"

        # Headers styling
        title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="0B2545", end_color="0B2545", fill_type="solid")
        ws.merge_cells("A1:H1")
        ws["A1"] = f"MANAK-AI Compliance Readiness Checklist — {metadata.get('product_name', 'Product')}"
        ws["A1"].font = title_font
        ws["A1"].fill = title_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        headers = ["ID", "Requirement", "Clause Reference", "Evidence Required", "Priority", "Status", "Owner", "Due Date"]
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="134074", end_color="134074", fill_type="solid")
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for row_idx, item in enumerate(checklist, 4):
            status = "COMPLETED" if item.get("completed") else "PENDING"
            ws.cell(row=row_idx, column=1, value=item.get("id", "")).border = thin_border
            ws.cell(row=row_idx, column=2, value=item.get("requirement", "")).border = thin_border
            ws.cell(row=row_idx, column=3, value=item.get("clause", "")).border = thin_border
            ws.cell(row=row_idx, column=4, value=item.get("evidence_needed", "")).border = thin_border
            ws.cell(row=row_idx, column=5, value=item.get("priority", "")).border = thin_border
            ws.cell(row=row_idx, column=6, value=status).border = thin_border
            ws.cell(row=row_idx, column=7, value=item.get("owner", "")).border = thin_border
            ws.cell(row=row_idx, column=8, value=item.get("due_date", "")).border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def export_checklist_pdf(self, checklist: List[Dict[str, Any]], metadata: Dict[str, Any]) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            textColor=colors.HexColor('#0B2545'),
            spaceAfter=10
        )
        subtitle_style = ParagraphStyle(
            'SubtitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#475569'),
            spaceAfter=15
        )

        elements.append(Paragraph("MANAK-AI — Pre-Audit Compliance Readiness Report", title_style))
        elements.append(Paragraph(f"Product: <b>{metadata.get('product_name', 'General Product')}</b> | Standard: <b>{metadata.get('standard_id', 'IS Standard')}</b> | Overall Readiness Score: <b>{metadata.get('overall_score', 85)}%</b>", subtitle_style))
        elements.append(Spacer(1, 10))

        table_data = [["ID", "Requirement", "Clause", "Priority", "Status", "Owner"]]
        for item in checklist:
            table_data.append([
                item.get("id", ""),
                Paragraph(item.get("requirement", "")[:70], styles['Normal']),
                item.get("clause", "")[:20],
                item.get("priority", ""),
                "DONE" if item.get("completed") else "PENDING",
                item.get("owner", "")[:15]
            ])

        t = Table(table_data, colWidths=[55, 230, 80, 55, 55, 65])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0B2545')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 9),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 8),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("<b>Notice:</b> This document was generated by MANAK-AI (Problem Statement SIH26107) for compliance preparation and pre-audit readiness. It does not replace official BIS audit reports.", styles['Italic']))

        doc.build(elements)
        return buffer.getvalue()

export_service = ExportService()
